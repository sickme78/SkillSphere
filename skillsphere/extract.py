"""Text extraction from uploaded PDF / PPTX / TXT learning material.

extract_text() returns the FULL readable text of the document (within sane safety
limits). Deciding how much of it goes to Gemini is material.py's job, so a long
document is no longer silently cut off after its first few pages.
"""
import os
import re

from pypdf import PdfReader
from pptx import Presentation

MAX_PAGES = 300           # safety limit for PDFs
MAX_RAW_CHARS = 600_000   # safety limit for everything we keep in memory
MAX_CHARS = 14_000        # kept for the JSON API: how much raw text it will accept


class ExtractError(ValueError):
    """Raised with a message that is safe to show to the learner."""


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = _from_pdf(path)
    elif ext == ".pptx":
        text = _from_pptx(path)
    elif ext == ".ppt":
        raise ExtractError("Old .ppt files are not supported. Open it in PowerPoint and "
                           "save it as .pptx (or export as PDF), then upload again.")
    elif ext == ".txt":
        with open(path, encoding="utf-8", errors="ignore") as fh:
            text = fh.read(MAX_RAW_CHARS)
    else:
        raise ExtractError("Unsupported file type: " + ext)
    return clean(text)[:MAX_RAW_CHARS]


def clean(text):
    """Normalise whitespace but keep paragraph breaks (chunking relies on them)."""
    text = text.replace("\r", "\n").replace("\x00", " ")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_pdf(path):
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractError("That PDF is password-protected. Remove the password and try again.")
    parts, size = [], 0
    for page in reader.pages[:MAX_PAGES]:
        page_text = page.extract_text() or ""
        parts.append(page_text)
        size += len(page_text)
        if size > MAX_RAW_CHARS:
            break
    return "\n\n".join(parts)


def _shape_text(shape):
    """Text from a shape, including grouped shapes and tables."""
    out = []
    if getattr(shape, "shapes", None) is not None:          # group shape
        for child in shape.shapes:
            out.extend(_shape_text(child))
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append(" | ".join(cell.text for cell in row.cells))
    return out


def _from_pptx(path):
    prs = Presentation(path)
    slides = []
    for slide in prs.slides:
        chunks = []
        for shape in slide.shapes:
            chunks.extend(_shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            chunks.append(slide.notes_slide.notes_text_frame.text)
        slides.append("\n".join(c for c in chunks if c and c.strip()))
    return "\n\n".join(s for s in slides if s.strip())
